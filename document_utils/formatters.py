import io
import os
import re
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from fpdf import FPDF


def sanitize_text(text: str) -> str:
    replacements = {
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2013": "-",
        "\u2014": "-",
        "\u00a0": " ",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)
    return text.strip()


def _split_content(text: str):
    lines = [line.strip() for line in sanitize_text(text).splitlines()]
    return [line for line in lines if line]


def format_docx(text: str, doc_type: str, logo_path: Optional[str] = None) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    if logo_path and os.path.exists(logo_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(logo_path, width=Inches(1.2))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(doc_type.upper())
    run.bold = True
    run.font.size = Pt(16)

    for line in _split_content(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(line)
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Generated with LegalEase - Draft for review").font.size = Pt(8)

    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


class LegalPDF(FPDF):
    def __init__(self, logo_path=None):
        super().__init__()
        self.logo_path = logo_path

    def header(self):
        if self.logo_path and os.path.exists(self.logo_path):
            try:
                self.image(self.logo_path, x=95, y=8, w=20)
                self.ln(16)
            except Exception:
                pass
        self.set_font("Helvetica", "B", 9)
        self.cell(0, 8, "LegalEase - Legal Document Draft", align="C")
        self.ln(8)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", "", 8)
        self.cell(0, 10, "Generated with LegalEase - Draft for review", align="C")


def format_pdf(text: str, doc_type: str, logo_path: Optional[str] = None) -> bytes:
    pdf = LegalPDF(logo_path)
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 15)
    pdf.cell(0, 10, sanitize_text(doc_type).upper(), align="C")
    pdf.ln(14)

    pdf.set_font("Helvetica", "", 11)

    for line in _split_content(text):
        safe = line.replace("•", "-")
        pdf.multi_cell(0, 6, safe)
        pdf.ln(2)

    return bytes(pdf.output())


def format_html_preview(text: str) -> str:
    escaped = (
        sanitize_text(text)
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )
    escaped = escaped.replace("\n", "<br>")
    return f'''
    <div style="
        background:#171717;
        color:#f2f2f2;
        padding:22px;
        border-radius:12px;
        line-height:1.65;
        max-height:600px;
        overflow-y:auto;
        border:1px solid #333;">
        {escaped}
    </div>
    '''
